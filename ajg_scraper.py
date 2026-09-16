import os, re, time, unicodedata, threading
from datetime import datetime
import openpyxl
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.support.ui import Select, WebDriverWait
from selenium.common.exceptions import (
    NoSuchElementException, TimeoutException, StaleElementReferenceException,
    UnexpectedAlertPresentException, NoAlertPresentException,
)
from selenium.webdriver.common.action_chains import ActionChains
import pyperclip
from util_log import escrever_log
from util_chrome import criar_driver
from paths import base_dir


def matar_processo_chrome():
    """Mata todos os processos chrome.exe e chromedriver.exe em execução."""
    try:
        from util_processo import matar_processo_chrome as _m
        _m()
    except Exception:
        pass

CAMINHO_RESULTADO = os.path.join(base_dir(), 'resultado.xlsx')
URL_LOGIN_AJG = ('https://ajg1.cjf.jus.br/aj/seguranca/efetuarloginintranet/'
                 'efetuarLoginIntranet_efetuarLogin.jsf')


class ScraperAJG:
    def __init__(self, debugger_address=None, vara=None, caminho_arquivo=None):
        self.logs = []
        self.log_event = threading.Event()
        self.login_event = threading.Event()
        self.stop_event = threading.Event()
        self.running = False
        self.done = False
        self.driver = None
        self.debugger_address = debugger_address
        self.vara = vara
        self.caminho_arquivo = caminho_arquivo  # Caminho do .xlsx uploadado (ou None para usar resultado.xlsx)

    def log(self, msg):
        ts = datetime.now().strftime('%H:%M:%S')
        entry = f'[{ts}] {msg}'
        self.logs.append(entry)
        self.log_event.set()
        print(entry)
        escrever_log(f'log_ajg_{datetime.now().strftime("%Y%m%d")}.txt', entry)

    def get_new_logs(self, last_idx):
        if last_idx < len(self.logs):
            return self.logs[last_idx:]
        return []

    @staticmethod
    def remover_acentos(texto):
        return ''.join(c for c in unicodedata.normalize('NFD', texto)
                       if unicodedata.category(c) != 'Mn')

    @staticmethod
    def digito(s):
        return re.sub(r'\D', '', s or '')

    def ler_resultado(self, caminho):
        wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
        ws = wb.active
        cabecalhos = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        resultados = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            dados_row = {}
            for col_idx, coluna in enumerate(cabecalhos):
                dados_row[coluna] = str(row[col_idx]) if row[col_idx] else ''
            if dados_row.get('nr_processo'):
                if dados_row.get('juntada', '').strip().upper() == 'OK':
                    dados_row['nome'] = dados_row.get('perito', '')
                    dados_row['data_servico'] = dados_row.get('data_nomeação', '')
                    dados_row['profissao'] = dados_row.get('profissao', '')
                    resultados.append(dados_row)
        wb.close()
        return resultados

    def esperar_ajg(self, driver):
        try:
            el = driver.find_element(By.ID, 'loadingDiv')
            return el.value_of_css_property('display') == 'none'
        except (NoSuchElementException, StaleElementReferenceException):
            return True

    def esperar_ajg_com_timeout(self, driver, timeout=120):
        WebDriverWait(driver, timeout).until(self.esperar_ajg)

    def handle_alert(self):
        """Aceita o alerta atual (se houver) e retorna o texto; senao None."""
        try:
            alerta = WebDriverWait(self.driver, 2).until(EC.alert_is_present())
            txt = alerta.text
            self.log(f'  ALERTA: {txt}')
            alerta.accept()
            return txt
        except TimeoutException:
            return None
        except NoAlertPresentException:
            return None

    def cancelar_e_voltar_inicio(self):
        """Clica cancelar (com confirmacao) e bt_inicio_sistema para ir ao proximo processo."""
        try:
            btn_cancelar = WebDriverWait(self.driver, 10).until(
                EC.element_to_be_clickable((By.ID, 'formAJIntranet:cancelar'))
            )
            btn_cancelar.click()
            time.sleep(1)
            self.handle_alert()  # confirmacao "Deseja cancelar a operacao..."
        except Exception as e:
            self.log(f'  Erro ao clicar cancelar: {e}')

        try:
            btn_inicio = WebDriverWait(self.driver, 10).until(
                EC.element_to_be_clickable((By.ID, 'formAJIntranet:bt_inicio_sistema'))
            )
            btn_inicio.click()
            self.esperar_ajg_com_timeout(self.driver)
            time.sleep(1)
        except Exception as e:
            self.log(f'  Erro ao voltar ao inicio: {e}')

    def navegar_para_nomeacao(self):
        """Abre a tela de nomeacao de profissionais (tambem ao voltar ao inicio
        entre processos)."""
        self.log('  PASSO: Procurando link "Nomeacao de Profissionais"...')
        try:
            el = WebDriverWait(self.driver, 10).until(
                EC.visibility_of_element_located((By.ID, 'link_tela_com_tooltip_1nomeacaodeprofissionais'))
            )
            el.click()
            self.log('  PASSO: Link "Nomeacao de Profissionais" clicado (por ID).')
        except TimeoutException:
            self.log('  PASSO: Link por ID nao encontrado, buscando por texto...')
            links = self.driver.find_elements(By.TAG_NAME, 'a')
            for l in links:
                lid = l.get_attribute('id') or ''
                ltext = l.text.strip()[:50]
                if 'nomeacao' in lid.lower() or 'nomeacao' in ltext.lower():
                    l.click()
                    self.log(f'  PASSO: Link clicado: id={lid} texto={ltext}')
                    break
            else:
                self.log('  ERRO: Nenhum link de nomeacao encontrado na pagina.')
                raise Exception('Link "Nomeacao de Profissionais" nao encontrado.')
        self.log('  PASSO: Aguardando carregamento da tela de nomeacao...')
        self.esperar_ajg_com_timeout(self.driver)
        time.sleep(1)
        self.log('  PASSO: Tela de nomeacao pronta.')

    def preencher_nomeacao(self, dados):
        """Novo -> PERITO -> profissao -> data -> polo passivo -> concBenef -> avancar."""
        btn_novo = WebDriverWait(self.driver, 10).until(
            EC.element_to_be_clickable((By.ID, 'formAJIntranet:novo'))
        )
        btn_novo.click()
        self.esperar_ajg_com_timeout(self.driver)
        time.sleep(1)

        Select(WebDriverWait(self.driver, 60).until(
            EC.visibility_of_element_located((By.ID, 'formAJIntranet:id_categoriasProfissionais'))
        )).select_by_visible_text('PERITO')

        profissoes = Select(WebDriverWait(self.driver, 60).until(
            EC.visibility_of_element_located((By.ID, 'formAJIntranet:id_profissoes'))
        ))
        for option in profissoes.options:
            if self.remover_acentos(option.text).lower() == self.remover_acentos(
                    dados.get('profissao', '')).lower():
                profissoes.select_by_visible_text(option.text)
                break
        else:
            raise Exception(f"Profissao '{dados.get('profissao')}' nao encontrada no AJG.")

        if dados.get('data_nomeacao'):
            campo_data = self.driver.find_element(By.ID, 'formAJIntranet:id_dataNomeacao')
            campo_data.clear()
            campo_data.send_keys(dados['data_nomeacao'])

        if 'INSS' in dados.get('polo_passivo', '').upper():
            self.driver.find_element(By.XPATH,
                '/html/body/form/div[3]/span/table/tbody/tr[7]/td[2]/fieldset/span/span/input[1]').click()
        else:
            self.driver.find_element(By.XPATH,
                '/html/body/form/div[3]/span/table/tbody/tr[7]/td[2]/fieldset/span/span/input[2]').click()

        self.driver.find_element(By.ID, 'formAJIntranet:concBenefAssistDeficBenefPrevIncLab').click()
        self.log('  Dados da nomeacao preenchidos.')
        self.driver.find_element(By.ID, 'formAJIntranet:avancar').click()

    def colar_numero_e_tab(self, dados):
        """Cola nr_processo + TAB. Retorna True se for preciso pular o processo."""
        campo = WebDriverWait(self.driver, 60).until(
            EC.visibility_of_element_located((By.ID, 'formAJIntranet:id_NumeroProcessoJudicial'))
        )
        campo.click()
        time.sleep(0.5)

        match_cnj = re.search(r'\d{7}-\d{2}\.\d{4}\.\d\.\d{2}\.\d{4}', dados.get('nr_processo', ''))
        nr_processo_ajg = match_cnj.group(0) if match_cnj else dados.get('nr_processo', '')

        pyperclip.copy(nr_processo_ajg)
        self.log(f'  Clipboard: {nr_processo_ajg}')
        ActionChains(self.driver).key_down(Keys.CONTROL).send_keys('v').key_up(Keys.CONTROL).perform()
        time.sleep(1)

        # alerta apos colar
        if self.handle_alert():
            return True

        try:
            valor_campo = campo.get_attribute('value')
        except StaleElementReferenceException:
            campo = WebDriverWait(self.driver, 10).until(
                EC.visibility_of_element_located((By.ID, 'formAJIntranet:id_NumeroProcessoJudicial'))
            )
            valor_campo = campo.get_attribute('value')
        self.log(f'  Valor no campo apos colar: {valor_campo}')

        if self.digito(valor_campo) != self.digito(nr_processo_ajg):
            self.log('  Colagem nao conferiu, pulando processo...')
            return True

        campo.send_keys(Keys.TAB)
        time.sleep(2)

        # alerta apos TAB
        if self.handle_alert():
            return True

        return False

    def selecionar_parte_ativa(self, dados):
        """Aguarda lista de partes, seleciona autor/requerente/exequente/impetrante
        (ou a unica parte), associa, salva parte_associada no Excel e avanca."""

        def _tem_parte_util(_driver):
            try:
                select_el = _driver.find_element(By.ID, 'formAJIntranet:selectOneListboxAssistido')
                for op in select_el.find_elements(By.TAG_NAME, 'option'):
                    txt = (op.text or '').strip().lower()
                    if txt and txt not in ('selecione', 'selecione...'):
                        return True
                return False
            except (NoSuchElementException, StaleElementReferenceException):
                return False

        try:
            WebDriverWait(self.driver, 20,
                          ignored_exceptions=(StaleElementReferenceException,)).until(_tem_parte_util)
        except TimeoutException:
            self.log('  Lista de partes NAO populou apos TAB. Tentando refresh + TAB...')
            self.driver.refresh()
            time.sleep(2)
            try:
                WebDriverWait(self.driver, 20,
                              ignored_exceptions=(StaleElementReferenceException,)).until(_tem_parte_util)
            except TimeoutException:
                self.log('  ERRO: lista de partes seguiu vazia. Pulando processo.')
                return False

        padrao_parte_ativa = re.compile(r'\b(autora?|requerente|exequente|impetrante)\b', re.IGNORECASE)
        autor_encontrado = False
        nome_parte_selecionada = None

        for _ in range(3):
            try:
                select_el = self.driver.find_element(By.ID, 'formAJIntranet:selectOneListboxAssistido')
                opcoes = select_el.find_elements(By.TAG_NAME, 'option')

                for op in opcoes:
                    texto = (op.text or '').strip()
                    if not texto or texto.lower() in ('selecione', 'selecione...'):
                        continue
                    if padrao_parte_ativa.search(texto):
                        WebDriverWait(self.driver, 10).until(EC.element_to_be_clickable(op))
                        op.click()
                        autor_encontrado = True
                        nome_parte_selecionada = texto
                        self.log(f'  Parte ativa selecionada: {texto}')
                        break

                if autor_encontrado:
                    break

                partes_reais = [op for op in opcoes
                                if (op.text or '').strip().lower() not in ('selecione', 'selecione...', '')]
                if len(partes_reais) == 1:
                    WebDriverWait(self.driver, 10).until(EC.element_to_be_clickable(partes_reais[0]))
                    partes_reais[0].click()
                    autor_encontrado = True
                    nome_parte_selecionada = partes_reais[0].text.strip()
                    self.log(f'  Unica parte real selecionada: {nome_parte_selecionada}')
                    break

            except StaleElementReferenceException:
                time.sleep(0.5)
                continue

        if not autor_encontrado:
            self.log('  ERRO: nao foi possivel selecionar parte ativa. Pulando processo.')
            return False

        # Associar parte
        self.driver.find_element(By.ID, 'formAJIntranet:id_0000001009').click()
        self.esperar_ajg_com_timeout(self.driver)
        time.sleep(1)
        self.log('  Parte associada com sucesso.')

        # Salvar parte_associada no resultado.xlsx
        try:
            wb = openpyxl.load_workbook(CAMINHO_RESULTADO)
            ws = wb.active
            cabecalhos = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
            col_parte = None
            for idx, col in enumerate(cabecalhos, 1):
                if col == 'parte_associada':
                    col_parte = idx
                    break
            if col_parte is None:
                col_parte = len(cabecalhos) + 1
                ws.cell(row=1, column=col_parte, value='parte_associada')

            nr_alvo = str(dados.get('nr_processo', '')).strip()
            for row in range(2, ws.max_row + 1):
                if ws.cell(row=row, column=1).value and str(ws.cell(row=row, column=1).value).strip() == nr_alvo:
                    ws.cell(row=row, column=col_parte, value=nome_parte_selecionada)
                    self.log(f'  Nome da parte salvo no Excel: {nome_parte_selecionada}')
                    break
            wb.save(CAMINHO_RESULTADO)
            wb.close()
        except Exception as e:
            self.log(f'  ERRO ao salvar parte no Excel (feche o arquivo se estiver aberto): {e}')

        # Avancar para honorarios
        self.driver.find_element(By.ID, 'formAJIntranet:avancar').click()
        self.esperar_ajg_com_timeout(self.driver)
        time.sleep(1)
        self.log('  Avancou para etapa de honorarios.')
        return True

    def preencher_honorarios_e_solicitacao(self, dados):
        """Honorarios -> busca perito -> valor -> concluir -> criar solicitacao -> final."""
        nome_perito = dados.get('nome', '')
        self.log(f'  [6/6] Preenchendo honorarios + solicitacao... (perito={nome_perito})')

        # Honorarios
        self.log('  [6/6] Selecionando PERITAS(OS)...')
        Select(self.driver.find_element(By.ID, 'formAJIntranet:id_honorarios')).select_by_visible_text('PERITAS(OS)')
        if dados.get('data_servico'):
            self.log(f'  [6/6] Preenchendo data servico: {dados["data_servico"]}')
            self.driver.find_element(By.ID, 'formAJIntranet:id_dataPrestServ').send_keys(dados['data_servico'])
        self.log('  [6/6] Clicando localPrestServ...')
        self.driver.find_element(By.ID, 'formAJIntranet:localPrestServ').click()
        time.sleep(1)

        # Salvar janela principal
        janela_principal = self.driver.current_window_handle
        janelas_antes = set(self.driver.window_handles)
        self.log(f'  [6/6] Janela principal: {janela_principal}')
        self.log(f'  [6/6] Janelas antes do click: {len(janelas_antes)}')

        # Clicar botao que abre popup
        self.log('  [6/6] Clicando botaoPesquisarProfissinal (abre popup)...')
        btn_pesq = WebDriverWait(self.driver, 30).until(
            EC.element_to_be_clickable((By.ID, 'formAJIntranet:botaoPesquisarProfissinal'))
        )
        self.driver.execute_script("arguments[0].click();", btn_pesq)
        time.sleep(3)

        # Detectar popup (nova janela)
        self.log('  [6/6] Aguardando popup...')
        janelas_depois = set(self.driver.window_handles)
        janelas_novas = janelas_depois - janelas_antes
        self.log(f'  [6/6] Janelas depois: {len(janelas_depois)}, novas: {len(janelas_novas)}')

        if janelas_novas:
            # Popup é uma nova janela
            handle_popup = janelas_novas.pop()
            self.log(f'  [6/6] Popup encontrado como nova janela: {handle_popup}')
            self.driver.switch_to.window(handle_popup)
            time.sleep(2)
        else:
            # Popup pode ser um iframe ou dialog no mesmo DOM
            self.log('  [6/6] Nenhuma nova janela. Buscando iframe ou dialog...')
            iframes = self.driver.find_elements(By.TAG_NAME, 'iframe')
            self.log(f'  [6/6] Iframes encontrados: {len(iframes)}')
            for idx, iframe in enumerate(iframes):
                iframe_id = iframe.get_attribute('id') or ''
                iframe_src = iframe.get_attribute('src') or ''
                self.log(f'    iframe[{idx}]: id={iframe_id} src={iframe_src[:80]}')
            # Tentar dialog/modal no DOM mesmo
            dialogs = self.driver.find_elements(By.CSS_SELECTOR, '[role="dialog"], .ui-dialog, .modal')
            self.log(f'  [6/6] Dialogs/modais: {len(dialogs)}')

        # Listar todas as janelas
        self.log(f'  [6/6] Janela atual: {self.driver.current_window_handle}')
        self.log(f'  [6/6] Todas janelas: {self.driver.window_handles}')

        # Buscar campo id_NomeProfissionalGuiaIndividual
        self.log('  [6/6] Buscando campo id_NomeProfissionalGuiaIndividual...')
        campo_nome = None
        seletores = [
            '#formAJIntranet\\:id_NomeProfissionalGuiaIndividual',
            'input[id*="NomeProfissionalGuia"]',
            'input[id*="GuiaIndividual"]',
            'input[id*="NomeProfissional"]',
        ]
        for sel in seletores:
            encontrados = self.driver.find_elements(By.CSS_SELECTOR, sel)
            self.log(f'  [6/6] Seletor "{sel}" -> {len(encontrados)} resultado(s)')
            if encontrados and not campo_nome:
                campo_nome = encontrados[0]
                self.log(f'  [6/6] Campo encontrado: id={campo_nome.get_attribute("id")} visible={campo_nome.is_displayed()}')

        # Fallback: buscar por todos inputs text na pagina atual
        if not campo_nome:
            self.log('  [6/6] Campo nao encontrado. Listando todos os inputs text da pagina atual...')
            all_text = self.driver.find_elements(By.CSS_SELECTOR, 'input[type="text"]')
            self.log(f'  [6/6] Total inputs text: {len(all_text)}')
            for idx, inp in enumerate(all_text):
                self.log(f'    [{idx}] id={inp.get_attribute("id")} value={inp.get_attribute("value") or ""} visible={inp.is_displayed()}')

        if campo_nome:
            self.log(f'  [6/6] Escrevendo "{nome_perito}" no campo...')
            # JS para forcar escrita
            self.driver.execute_script("""
                var el = arguments[0];
                el.removeAttribute('readonly');
                el.removeAttribute('disabled');
                el.focus();
            """, campo_nome)
            time.sleep(0.3)
            ActionChains(self.driver).click(campo_nome).perform()
            time.sleep(0.3)
            campo_nome.send_keys(Keys.CONTROL + 'a')
            time.sleep(0.2)
            campo_nome.send_keys(Keys.DELETE)
            time.sleep(0.2)
            campo_nome.send_keys(nome_perito)
            time.sleep(0.5)
            valor = campo_nome.get_attribute('value') or ''
            self.log(f'  [6/6] Valor no campo: "{valor}"')

            # Se vazio, forcar via JS
            if not valor:
                self.log('  [6/6] Valor vazio, forçando via JS...')
                self.driver.execute_script("arguments[0].value = arguments[1];", campo_nome, nome_perito)
                time.sleep(0.3)
                valor = campo_nome.get_attribute('value') or ''
                self.log(f'  [6/6] Valor apos JS: "{valor}"')

            # Clicar botao pesquisar no popup (id=FormAJIntranet:pesquisar)
            self.log('  [6/6] Buscando botao pesquisar no popup...')
            btn_pesq_popup = None

            # 1. Buscar por ID exato (com variações de case)
            for id_tentativa in ['FormAJIntranet:pesquisar', 'formAJIntranet:pesquisar']:
                encontrados = self.driver.find_elements(By.ID, id_tentativa)
                self.log(f'  [6/6] ID "{id_tentativa}" -> {len(encontrados)} resultado(s)')
                for en in encontrados:
                    self.log(f'    tag={en.tag_name} type={en.get_attribute("type")} visible={en.is_displayed()}')
                    if en.is_displayed() and not btn_pesq_popup:
                        btn_pesq_popup = en

            # 2. Buscar por input/image com "pesquisar" no ID
            if not btn_pesq_popup:
                candidatos = self.driver.find_elements(By.CSS_SELECTOR,
                    "input[id*='esquisar'], input[id*='esq'], button[id*='esquisar'], input[type='image'][id*='esquisar']"
                )
                self.log(f'  [6/6] Busca por ID parcial: {len(candidatos)} resultado(s)')
                for idx, btn in enumerate(candidatos):
                    vis = btn.is_displayed()
                    self.log(f'    [{idx}] id={btn.get_attribute("id")} type={btn.get_attribute("type")} visible={vis}')
                    if vis and not btn_pesq_popup:
                        btn_pesq_popup = btn

            # 3. Buscar por xpath: input com value/text "pesquisar"
            if not btn_pesq_popup:
                candidatos = self.driver.find_elements(By.XPATH,
                    "//*[contains(translate(@value,'ABCDEFGHIJKLMNOPQRSTUVWXYZ','abcdefghijklmnopqrstuvwxyz'),'pesquisar') or "
                    "contains(translate(text(),'ABCDEFGHIJKLMNOPQRSTUVWXYZ','abcdefghijklmnopqrstuvwxyz'),'pesquisar')]"
                )
                self.log(f'  [6/6] Busca por texto: {len(candidatos)} resultado(s)')
                for idx, btn in enumerate(candidatos):
                    vis = btn.is_displayed()
                    self.log(f'    [{idx}] tag={btn.tag_name} id={btn.get_attribute("id")} type={btn.get_attribute("type")} visible={vis}')
                    if vis and not btn_pesq_popup:
                        btn_pesq_popup = btn

            if btn_pesq_popup:
                self.log(f'  [6/6] Clicando pesquisar: id={btn_pesq_popup.get_attribute("id")}...')
                self.driver.execute_script("arguments[0].click();", btn_pesq_popup)
                time.sleep(4)
            else:
                self.log('  [6/6] Botao nao encontrado, tentando Enter...')
                campo_nome.send_keys(Keys.RETURN)
                time.sleep(4)

            # Listar radios disponiveis
            self.log('  [6/6] Buscando radios na pagina do popup...')
            radios = self.driver.find_elements(By.CSS_SELECTOR, "input[type='radio']")
            self.log(f'  [6/6] Total de radios: {len(radios)}')
            for idx, radio in enumerate(radios):
                onclick = radio.get_attribute('onclick') or ''
                self.log(f'    radio[{idx}]: visible={radio.is_displayed()} onclick={onclick[:150]}')

            # Buscar perito nos radios
            nome_norm = self.remover_acentos(nome_perito).lower()
            encontrado = False
            for radio in radios:
                onclick = radio.get_attribute('onclick') or ''
                if nome_norm in self.remover_acentos(onclick).lower():
                    self.driver.execute_script("arguments[0].click();", radio)
                    encontrado = True
                    self.log(f'  Perito selecionado: {nome_perito}')
                    break

            # Busca parcial
            if not encontrado:
                self.log('  [6/6] Busca exata falhou. Tentando parcial...')
                for radio in radios:
                    onclick = radio.get_attribute('onclick') or ''
                    texto = self.remover_acentos(onclick).lower()
                    for palavra in nome_norm.split():
                        if len(palavra) > 2 and palavra in texto:
                            self.driver.execute_script("arguments[0].click();", radio)
                            self.log(f'  Perito selecionado por parcial: {nome_perito}')
                            encontrado = True
                            break
                    if encontrado:
                        break

            if not encontrado:
                self.log(f'  ERRO: Perito "{nome_perito}" nao encontrado nos radios. Pulando.')
                if len(self.driver.window_handles) > 1:
                    self.driver.close()
                self.driver.switch_to.window(janela_principal)
                return False

            # Clicar botao CONCLUIR no popup (id=FormAJIntranet:pesquisar type=image)
            time.sleep(1)
            self.log('  [6/6] Clicando concluir no popup...')
            btn_concluir = None
            for id_tentativa in ['FormAJIntranet:pesquisar', 'formAJIntranet:pesquisar']:
                encontrados = self.driver.find_elements(By.ID, id_tentativa)
                for en in encontrados:
                    tag = en.tag_name
                    tipo = en.get_attribute('type') or ''
                    vis = en.is_displayed()
                    self.log(f'    ID "{id_tentativa}": tag={tag} type={tipo} visible={vis}')
                    if vis and not btn_concluir:
                        btn_concluir = en

            if not btn_concluir:
                # Buscar input type=image com "concluir" no src
                imgs = self.driver.find_elements(By.CSS_SELECTOR, "input[type='image']")
                self.log(f'  [6/6] Inputs image encontrados: {len(imgs)}')
                for idx, img in enumerate(imgs):
                    src = img.get_attribute('src') or ''
                    vis = img.is_displayed()
                    self.log(f'    img[{idx}]: id={img.get_attribute("id")} src={src[:80]} visible={vis}')
                    if vis and ('concluir' in src.lower() or 'pesquis' in src.lower()):
                        btn_concluir = img
                        break

            if btn_concluir:
                self.log('  [6/6] Clicando concluir...')
                self.driver.execute_script("arguments[0].click();", btn_concluir)
                time.sleep(3)
            else:
                self.log('  [6/6] Botao concluir nao encontrado no popup')

        else:
            self.log(f'  ERRO: Campo de busca nao encontrado. Pulando processo.')
            if len(self.driver.window_handles) > 1:
                self.driver.close()
            self.driver.switch_to.window(janela_principal)
            return False

        # Voltar para janela principal
        self.log('  [6/6] Fechando popup e voltando para janela principal...')
        if len(self.driver.window_handles) > 1:
            self.driver.close()
        self.driver.switch_to.window(janela_principal)
        time.sleep(2)

        # Avancar para valor
        self.log('  [6/6] Avancando para tela de valor estimado...')
        avancar = self.driver.find_element(By.ID, 'formAJIntranet:avancar')
        self.driver.execute_script("arguments[0].click();", avancar)
        self.esperar_ajg_com_timeout(self.driver)
        time.sleep(2)

        # Valor estimado honorarios
        self.log(f'  [6/6] Buscando campo valor estimado... (valor={dados.get("valor")})')
        campo_valor = WebDriverWait(self.driver, 60).until(
            EC.visibility_of_element_located((By.ID, 'formAJIntranet:id_valorEstimadoHonorarios'))
        )
        self.log(f'  [6/6] Campo valor encontrado. Preenchendo...')
        if dados.get('valor'):
            campo_valor.click()
            time.sleep(0.3)
            campo_valor.send_keys(str(dados['valor']))
            time.sleep(0.5)
            valor_escrito = campo_valor.get_attribute('value') or ''
            self.log(f'  [6/6] Valor escrito no campo: "{valor_escrito}"')
        self.log('  Valor estimado preenchido.')

        # Clicar concluir
        self.log('  [6/6] Clicando concluir...')
        concluir = WebDriverWait(self.driver, 60).until(
            EC.element_to_be_clickable((By.ID, 'formAJIntranet:concluir'))
        )
        self.driver.execute_script("arguments[0].click();", concluir)
        self.esperar_ajg_com_timeout(self.driver)
        time.sleep(1)
        self.log('  Concluido.')

        # Criar solicitacao
        self.log('  [6/6] Criando solicitacao de pagamento...')
        WebDriverWait(self.driver, 60).until(
            EC.element_to_be_clickable((By.ID, 'formAJIntranet:criarSolicitacao'))
        )
        self.driver.execute_script("arguments[0].click();", self.driver.find_element(By.ID, 'formAJIntranet:criarSolicitacao'))
        self.esperar_ajg_com_timeout(self.driver)
        self.log('  Criei solicitacao de pagamento.')

        # Dados finais da solicitacao
        WebDriverWait(self.driver, 60).until(
            EC.visibility_of_element_located((By.ID, 'formAJIntranet:dataPrestacao'))
        ).send_keys(dados.get('data_servico', ''))

        self.driver.find_element(By.ID, 'copiar').click()
        self.driver.find_element(By.ID, 'formAJIntranet:motivos2').click()
        self.driver.find_element(By.ID, 'formAJIntranet:decisaoFundamentada').send_keys('.')

        WebDriverWait(self.driver, 60).until(
            EC.visibility_of_element_located((By.ID, 'formAJIntranet:concluir'))
        ).click()
        self.esperar_ajg_com_timeout(self.driver)
        time.sleep(1)
        self.log('  Cliquei em concluir (solicitacao criada).')

        # Voltar ao inicio para o proximo processo
        WebDriverWait(self.driver, 120).until(self.esperar_ajg)
        self.driver.find_element(By.ID, 'formAJIntranet:bt_inicio_sistema').click()
        self.esperar_ajg_com_timeout(self.driver)
        time.sleep(1)

    def run(self):
        self.running = True
        self.done = False

        try:
            # Usa arquivo uploadado (caminho_arquivo) ou fallback para resultado.xlsx do PJe
            caminho = self.caminho_arquivo if self.caminho_arquivo and os.path.exists(self.caminho_arquivo) else CAMINHO_RESULTADO
            nome_arquivo = os.path.basename(caminho)
            self.log(f'ETAPA 1: Lendo {nome_arquivo}...')
            if not os.path.exists(caminho):
                self.log(f'ERRO: {nome_arquivo} nao encontrado.')
                return

            resultados = self.ler_resultado(caminho)

            if self.vara:
                resultados = [r for r in resultados if r.get('vara', '').strip() == self.vara]
                self.log(f'Filtrando por vara "{self.vara}": {len(resultados)} processo(s)')

            self.log(f'Processos com juntada=OK: {len(resultados)}')

            if not resultados:
                self.log('Nenhum processo com juntada=OK.')
                return

            for idx, r in enumerate(resultados, 1):
                self.log(f'  [{idx}] {r.get("nr_processo")} | {r.get("nome")} | {r.get("profissao")} | {r.get("data_servico")} | {r.get("valor")}')

            self.log('ETAPA 2: Abrindo Chrome AJG...')
            try:
                # Com debugger_address=None, abre Chrome novo (mais rapido)
                self.log('Abrindo Chrome AJG (nova instância)...')
                self.driver = criar_driver()
                self.driver.get(URL_LOGIN_AJG)
                self.log('Chrome AJG aberto com sucesso.')
            except Exception as e:
                self.log(f'Erro ao abrir Chrome AJG: {e}')
                return

            self.log('ETAPA 3: Aguardando login manual...')
            self.log('Faca login e clique em "Ja loguei no AJG".')
            self.login_event.wait()
            if self.stop_event.is_set():
                self.log('Parado pelo usuario.')
                return
            self.log('Login AJG confirmado!')

            self.log('ETAPA 4: Navegando para Nomeacao de Profissionais...')
            self.driver.switch_to.window(self.driver.window_handles[0])
            time.sleep(2)
            self.navegar_para_nomeacao()
            self.log('Tela de Nomeacao de Profissionais carregada!')

            self.log(f'ETAPA 5: Processando {len(resultados)} processo(s)...')
            for i, dados in enumerate(resultados, 1):
                if self.stop_event.is_set():
                    self.log('Execucao interrompida pelo usuario.')
                    break

                nr_processo = dados.get('nr_processo', '').strip()
                self.log(f'--- Processo {i}/{len(resultados)}: {nr_processo} ---')

                try:
                    if i > 1:
                        self.log('  Voltando para Nomeacao de Profissionais...')
                        self.navegar_para_nomeacao()

                    self.log('  [1/6] Novo -> PERITO -> profissao -> data -> avancar...')
                    try:
                        self.preencher_nomeacao(dados)
                    except UnexpectedAlertPresentException:
                        pass

                    self.log('  [2/6] Verificando alerta apos avancar...')
                    alerta = self.handle_alert()
                    if alerta and ('paga' in alerta.lower() or 'nomea' in alerta.lower()):
                        self.log('  ALERTA: nomeacao paga. Cancelando.')
                        self.cancelar_e_voltar_inicio()
                        continue

                    self.log('  [3/6] Aguardando loading apos avancar...')
                    try:
                        WebDriverWait(self.driver, 120).until(self.esperar_ajg)
                    except TimeoutException:
                        self.log('  Timeout no loading. Recarregando...')
                        self.driver.refresh()
                        time.sleep(2)

                    self.log('  [4/6] Colando numero do processo + TAB...')
                    if self.colar_numero_e_tab(dados):
                        self.log('  ALERTA: problema ao colar numero. Cancelando.')
                        self.cancelar_e_voltar_inicio()
                        continue

                    self.log('  [5/6] Selecionando parte ativa...')
                    if not self.selecionar_parte_ativa(dados):
                        self.cancelar_e_voltar_inicio()
                        continue

                    self.preencher_honorarios_e_solicitacao(dados)
                    self.log(f'  PROCESSO {nr_processo} CONCLUIDO!')

                except Exception as e:
                    import traceback
                    self.log(f'  ERRO no processo {nr_processo}: {e}')
                    self.log(traceback.format_exc())
                    try:
                        self.cancelar_e_voltar_inicio()
                    except Exception:
                        pass
                    continue

            self.log('=== EXECUCAO FINALIZADA ===')

        except Exception as e:
            import traceback
            self.log(f'ERRO GERAL: {type(e).__name__}: {e}')
            self.log(traceback.format_exc())
        finally:
            self.running = False
            self.done = True

    def confirmar_login(self):
        self.login_event.set()

    def fechar_navegador(self):
        if self.driver:
            try:
                self.driver.quit()
                self.log('Navegador fechado.')
            except:
                pass

    def parar(self):
        self.stop_event.set()
        self.login_event.set()
        if self.driver:
            try:
                self.driver.quit()
            except:
                pass
        self.driver = None
        # Mata processos Chrome/ChromeDriver em segundo plano
        matar_processo_chrome()
